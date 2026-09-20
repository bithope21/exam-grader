import json

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook
from test_workflow import prepare

from exam_grader.exporting import (
    checked_filename,
    checked_image,
    export_results,
    human_readable_name,
)
from exam_grader.imaging import reference_image


def test_export_consistent_versioned_and_keeps_originals(tmp_path):
    flow, exam, key_source, source = prepare(tmp_path)
    key = flow.approve_key(exam.id, ["A", "B", "C"], key_source["id"])
    flow.review(source["id"], "01", ["A", "blank", "C"], key["id"])
    first = export_results(flow, exam.id)
    snapshot = json.loads((first / "_system" / "results.json").read_text())
    workbook = load_workbook(first / "scores.xlsx")
    assert workbook.active.cell(2, 1).value == 1
    assert workbook.active.cell(2, 1).data_type == "n"
    assert workbook.active.cell(2, 2).value == snapshot["results"][0]["score"] == 2
    assert (first / snapshot["results"][0]["checked_image"]).is_file()
    second = export_results(flow, exam.id)
    assert second != first
    assert first.is_dir()
    assert flow.database.exists()
    assert (
        len(
            __import__("sqlite3")
            .connect(flow.database)
            .execute("select * from export_runs where exam_id=?", (exam.id,))
            .fetchall()
        )
        == 2
    )
    assert snapshot["export"]["human_readable_name"]
    assert snapshot["export"]["color_legend"]["correct"] == "green"
    assert snapshot["results"][0]["checked_provenance"]["encoding"]["format"] == "JPEG"
    assert first.parent.parent.name == snapshot["export"]["human_readable_name"]
    assert {p.name for p in first.iterdir()} == {"checked", "scores.xlsx", "_system"}
    assert (first / "checked" / "เลขที่-01.jpg").exists()


def test_human_readable_name_sanitizes_cross_platform_fields():
    assert human_readable_name(
        {
            "academic_year": "2569",
            "grade": "M/6",
            "room": "1?",
            "subject": "คณิต",
            "name": "สอบ:*",
            "question_count": 30,
        }
    ).endswith("_30q")
    assert "/" not in human_readable_name(
        {
            "academic_year": "a/b",
            "grade": "g",
            "room": "r",
            "subject": "s",
            "name": "n",
            "question_count": 1,
        }
    )
    assert len(human_readable_name({"name": "คณิตศาสตร์" * 100}).encode("utf-8")) < 255


def test_filename_collision_and_invalid_number_never_expose_ids_or_paths():
    used: set[str] = set()
    assert checked_filename("02", used) == "เลขที่-02.jpg"
    assert checked_filename("๒", used) == "เลขที่-02_02.jpg"
    assert checked_filename("../../bad", used) == "ไม่ทราบเลขที่-01.jpg"
    assert checked_filename(None, used) == "ไม่ทราบเลขที่-02.jpg"


@pytest.mark.parametrize("maximum", [1, 30, 60])
def test_score_occupies_printed_score_box_and_never_number_line(maximum):
    image = reference_image()
    source = cv2.imencode(".png", image)[1].tobytes()
    result = {
        "answers": ["A"] * maximum,
        "score": maximum,
        "max": maximum,
        "student_number": "1",
        "detection": {"registration": {"matrix": np.eye(3).tolist()}},
    }
    pixels = cv2.imdecode(
        np.frombuffer(checked_image(source, result, {"answers": result["answers"]}), np.uint8),
        cv2.IMREAD_COLOR,
    )
    assert pixels is not None
    blue = (pixels[:, :, 0] > 120) & (pixels[:, :, 1] < 150) & (pixels[:, :, 2] < 100)
    # Independent coordinates from the inspected printed reference, deliberately
    # not loaded from the implementation's score_roi.
    assert blue[55:117, 713:808].sum() > 100
    assert not blue[70:130, 590:700].any()
    assert not blue[15:48, 700:820].any()  # printed score title remains untouched


def test_identity_correction_regenerates_filename_and_excel_keeps_prior_run(tmp_path):
    flow, exam, key_source, source = prepare(tmp_path)
    key = flow.approve_key(exam.id, ["A", "B", "C"], key_source["id"])
    flow.review(source["id"], "1", ["A", "B", "C"], key["id"])
    first = export_results(flow, exam.id)
    original = (first / "_system/results.json").read_bytes()
    flow.review(source["id"], "12", ["A", "B", "C"], key["id"])
    second = export_results(flow, exam.id)
    assert (second / "checked/เลขที่-12.jpg").exists()
    assert load_workbook(second / "scores.xlsx").active.cell(2, 1).value == 12
    assert (first / "_system/results.json").read_bytes() == original


def test_checked_renderer_marks_canonical_cells_and_score_without_source_mutation():
    image = reference_image()
    source = cv2.imencode(".png", image)[1].tobytes()
    key = {"answers": ["A"]}
    result = {
        "answers": ["A"],
        "score": 1,
        "max": 1,
        "student_number": "1",
        "detection": {"registration": {"matrix": np.eye(3).tolist()}, "template_id": "test"},
    }
    rendered = checked_image(source, result, key)
    decoded = cv2.imdecode(np.frombuffer(rendered, np.uint8), cv2.IMREAD_COLOR)
    assert decoded is not None
    assert decoded.shape == image.shape
    green = (decoded[:, :, 1] > 120) & (decoded[:, :, 0] < 100) & (decoded[:, :, 2] < 100)
    assert int(green.sum()) > 20
    assert source == cv2.imencode(".png", image)[1].tobytes()


def test_checked_renderer_uses_red_for_wrong_and_distinct_non_answer_colors():
    image = reference_image()
    source = cv2.imencode(".png", image)[1].tobytes()
    result = {
        "answers": ["B", "blank", "multiple", "boundary_cross"],
        "score": 0,
        "max": 4,
        "student_number": "1",
        "detection": {
            "registration": {"matrix": np.eye(3).tolist()},
            "answers": [
                {"selected": ["B"]},
                {"selected": []},
                {"selected": ["A", "C"]},
                {"selected": ["B", "C"]},
            ],
        },
    }
    decoded = cv2.imdecode(
        np.frombuffer(checked_image(source, result, {"answers": ["A"] * 4}), np.uint8),
        cv2.IMREAD_COLOR,
    )
    assert decoded is not None
    red = (decoded[:, :, 2] > 140) & (decoded[:, :, 1] < 120) & (decoded[:, :, 0] < 120)
    orange = (decoded[:, :, 2] > 150) & (decoded[:, :, 1] > 70) & (decoded[:, :, 0] < 100)
    purple = (decoded[:, :, 2] > 100) & (decoded[:, :, 0] > 100) & (decoded[:, :, 1] < 130)
    assert int(red.sum()) > 20
    assert int(orange.sum()) > 20
    assert int(purple.sum()) > 20


def test_tampered_source_prevents_export(tmp_path):
    flow, exam, key_source, source = prepare(tmp_path)
    key = flow.approve_key(exam.id, ["A", "B", "C"], key_source["id"])
    flow.review(source["id"], "1", ["A", "B", "C"], key["id"])
    (flow.database.parent / source["relative_path"]).write_bytes(b"changed")
    with pytest.raises(ValueError):
        export_results(flow, exam.id)
    assert not list((flow.database.parent / "results" / exam.id).glob("run-*"))


def test_export_accepts_per_exam_output_root_and_rejects_file(tmp_path):
    flow, exam, key_source, source = prepare(tmp_path)
    key = flow.approve_key(exam.id, ["A", "B", "C"], key_source["id"])
    flow.review(source["id"], "1", ["A", "B", "C"], key["id"])
    output_root = tmp_path / "chosen-results"
    result = export_results(flow, exam.id, output_root=output_root)
    assert result.parent.parent.parent == output_root
    assert exam.id not in str(result)
    blocked = tmp_path / "not-a-folder"
    blocked.write_text("keep")
    with pytest.raises(OSError):
        export_results(flow, exam.id, output_root=blocked)


def test_custom_template_checked_overlay_alignment(tmp_path):
    """Verify custom template geometry is used for checked overlays, not Default #1."""
    import dataclasses
    import json

    import cv2
    import numpy as np

    from exam_grader.app import initialize
    from exam_grader.domain import ExamDetails
    from exam_grader.exporting import export_results
    from exam_grader.imports import ImportService
    from exam_grader.template_manager import load_builtin_template, save_custom_reference_image
    from exam_grader.workflow import Workflow

    app_data = tmp_path / "app_data"
    app = initialize(app_data)
    t3 = load_builtin_template("default-3")
    # Save reference image to content-addressed storage for custom template
    from exam_grader.template_manager import get_reference_image
    ref_img = get_reference_image(t3)
    ref_bytes = cv2.imencode(".png", ref_img)[1].tobytes()
    ref_sha = save_custom_reference_image(app_data, ref_bytes)

    # Create a custom template in the database with custom ID and 1200x1720 geometry
    custom_tmpl = dataclasses.replace(
        t3,
        template_id="custom-test-t3-001",
        name="Custom Default 3 Test",
        kind="custom",
        reference_sha256=ref_sha,
    )
    app.exams.save_template(custom_tmpl)

    # 2. Create exam with custom template
    exam = app.exams.create(
        ExamDetails(
            name="Custom Template Exam",
            academic_year="2569",
            subject="ศิลปะ",
            grade="ป.4",
            room="1",
            question_count=30,
            template_id=custom_tmpl.template_id,
            template_version=custom_tmpl.version,
        )
    )

    importer = ImportService(app.exams.path)
    flow = Workflow(app.exams.path)

    # Use default 3 reference image as dummy test sheet
    from exam_grader.template_manager import get_reference_image
    ref_img = get_reference_image(custom_tmpl, app_data_dir=tmp_path / "app_data")
    sheet_bytes = cv2.imencode(".png", ref_img)[1].tobytes()

    key_file = tmp_path / "key.png"
    key_file.write_bytes(sheet_bytes)

    student_img = ref_img.copy()
    student_img[0, 0] = [200, 200, 200]
    student_file = tmp_path / "student.png"
    student_file.write_bytes(cv2.imencode(".png", student_img)[1].tobytes())

    k_src = importer.import_file(exam.id, key_file, "key")
    s_src = importer.import_file(exam.id, student_file, "student")

    # Simulate detection with identity matrix on the custom template size
    dummy_detection = {
        "registration": {"matrix": np.eye(3).tolist(), "inliers": 100},
        "answers": [{"question": i, "selected": ["A"]} for i in range(1, 31)],
    }
    flow.save_detection(k_src["id"], dummy_detection)
    flow.save_detection(s_src["id"], dummy_detection)

    approved_key = flow.approve_key(exam.id, ["A"] * 30, k_src["id"])
    flow.review(s_src["id"], "13", ["A"] * 30, approved_key["id"])

    # 3. Export results
    export_dir = export_results(flow, exam.id)
    assert export_dir.is_dir()

    results_json_path = export_dir / "_system" / "results.json"
    assert results_json_path.exists()
    snapshot = json.loads(results_json_path.read_text(encoding="utf-8"))

    # Invariant checks:
    assert snapshot["export"]["template"]["template_id"] == custom_tmpl.template_id
    res = snapshot["results"][0]
    prov = res["checked_provenance"]
    assert prov["template_id"] == custom_tmpl.template_id
    assert prov["render_template"]["canonical_width"] == 1200
    assert prov["render_template"]["canonical_height"] == 1720

    # Inspect rendered checked image: verify marks were drawn inside the custom template's cells
    checked_path = export_dir / res["checked_image"]
    assert checked_path.exists()
    checked_img = cv2.imread(str(checked_path))
    assert checked_img is not None
    assert checked_img.shape[0] == 1720 and checked_img.shape[1] == 1200

    # Cell 1 choice A in custom template (Default #3):
    # Block 0: col_boundaries[0]=205, col_boundaries[1]=241, row_boundaries[0]=603
    # Check green color in block 0 cell 1
    block0 = custom_tmpl.answer_blocks[0]
    x1 = block0.col_boundaries[0]
    x2 = block0.col_boundaries[1]
    y1 = block0.row_boundaries[0]
    y2 = block0.row_boundaries[1]
    cell_crop = checked_img[y1:y2, x1:x2]
    # Green mark should be present in this cell
    green = (cell_crop[:, :, 1] > 140) & (cell_crop[:, :, 2] < 100) & (cell_crop[:, :, 0] < 100)
    assert green.sum() > 20, "Green checkmark must be rendered inside custom template cell (Q1/A)"
