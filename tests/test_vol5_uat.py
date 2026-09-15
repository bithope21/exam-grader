"""End-to-end Real-Fixture UAT for vol.5 (Default #2 Pink Sheets) and Permanent Purge."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from exam_grader.app import initialize
from exam_grader.domain import ExamDetails
from exam_grader.exporting import export_results
from exam_grader.identity import observe
from exam_grader.imaging import analyze
from exam_grader.imports import ImportService
from exam_grader.template_manager import load_builtin_template
from exam_grader.workflow import Workflow


def test_vol5_real_fixtures_automatic_omr_and_grading():
    vol5_dir = Path("tests/fixtures/real/vol.5")
    if not vol5_dir.exists():
        pytest.skip("vol.5 fixtures not present")

    t2 = load_builtin_template("default-2")

    # 1. Key must be automatically read with 30 answers
    key_bytes = (vol5_dir / "key.JPG").read_bytes()
    key_res = analyze(key_bytes, template_def=t2)
    assert key_res["registration"]["table_coverage"] >= 0.995
    key_answers = [a["selected"][0] if a["selected"] else None for a in key_res["answers"][:30]]
    # Expected key as confirmed by teacher:
    # 1:B, 2:B, 3:B, 4:A, 5:D, 6:C, 7:B, 8:B, 9:C, 10:B, 11:A, 12:D, 13:C, 14:D, 15:A,
    # 16:B, 17:B, 18:D, 19:B, 20:C, 21:A, 22:B, 23:C, 24:B, 25:A, 26:D, 27:A, 28:B, 29:D, 30:B
    expected_key = [
        "B",
        "B",
        "B",
        "A",
        "D",
        "C",
        "B",
        "B",
        "C",
        "B",
        "A",
        "D",
        "C",
        "D",
        "A",
        "B",
        "B",
        "D",
        "B",
        "C",
        "A",
        "B",
        "C",
        "B",
        "A",
        "D",
        "A",
        "B",
        "D",
        "B",
    ]
    assert key_answers == expected_key, f"Key mismatch: {key_answers} vs {expected_key}"

    # 2. Student 1 (IMG_0865.jpg) must remain available for teacher review.
    s1_bytes = (vol5_dir / "IMG_0865.jpg").read_bytes()
    s1_res = analyze(s1_bytes, template_def=t2)
    assert s1_res["registration"]["table_coverage"] >= 0.995
    s1_answers = [a["selected"][0] if a["selected"] else None for a in s1_res["answers"][:30]]
    # Grade student 1 against key:
    s1_score = sum(1 for a, k in zip(s1_answers, expected_key) if a == k)
    assert s1_score == 11, f"Student 1 score expected 11, got {s1_score}"

    # Verify student 1 identity observation
    s1_identity = observe(s1_bytes, s1_res["registration"]["matrix"], template_def=t2)
    assert s1_identity["candidate"] == "1" or "1" in s1_identity["candidates"]
    assert s1_identity["requires_review"] is True


def test_vol5_full_workflow_export_matches_ground_truth(tmp_path):
    vol5_dir = Path("tests/fixtures/real/vol.5")
    if not vol5_dir.exists():
        pytest.skip("vol.5 fixtures not present")

    app = initialize(tmp_path / "app_data")
    t2 = load_builtin_template("default-2")

    exam = app.exams.create(
        ExamDetails(
            name="คอมพิวเตอร์ ป.2/1 กระดาษชมพู",
            academic_year="2569",
            grade="ป.2",
            room="1",
            subject="com",
            question_count=30,
            template_id="default-2",
            template_version=1,
        )
    )

    importer = ImportService(app.exams.path)
    flow = Workflow(app.exams.path)

    # Import and approve key
    key_record = importer.import_file(exam.id, vol5_dir / "key.JPG", purpose="key")
    key_res = analyze(importer.verified_bytes(key_record), template_def=t2)
    key_answers = [a["selected"][0] for a in key_res["answers"][:30]]
    key = flow.approve_key(exam.id, key_answers, key_record["id"])

    # Import student sheets
    student_files = {
        "IMG_0865.jpg": "1",
        "IMG_0863.JPG": "2",
        "IMG_0864.jpg": "3",
        "IMG_0862.JPG": "4",
    }
    for fname, expected_no in student_files.items():
        src = importer.import_file(exam.id, vol5_dir / fname, purpose="student")
        res = analyze(importer.verified_bytes(src), template_def=t2)
        flow.save_detection(src["id"], res)
        obs = observe(importer.verified_bytes(src), res["registration"]["matrix"], template_def=t2)
        candidates = obs.get("candidates") or []
        assert obs.get("candidate") == expected_no or expected_no in candidates, (
            f"Expected {expected_no} to remain available for review; "
            f"primary={obs.get('candidate')}, candidates={candidates} for {fname}"
        )
        assert obs["requires_review"] is True
        # Fixture truth is an explicit teacher confirmation; an uncalibrated
        # candidate is never adopted as authoritative here.
        detected_no = expected_no
        flow.review(
            src["id"],
            detected_no,
            [a["selected"][0] if a["selected"] else "blank" for a in res["answers"][:30]],
            key["id"],
        )

    # Export results
    export_dir = export_results(flow, exam.id, output_root=tmp_path / "results")
    assert export_dir.is_dir()
    excel_file = export_dir / "scores.xlsx"
    assert excel_file.is_file()

    wb = load_workbook(excel_file)
    scores_sheet = wb["Scores"]
    data = {}
    for r in range(2, scores_sheet.max_row + 1):
        no = str(scores_sheet.cell(r, 1).value)
        score = scores_sheet.cell(r, 2).value
        data[no] = score

    # Compare with ground truth:
    # Student 1: 11
    # Student 2: 6
    # Student 3: 6
    # Student 4: 29
    assert data["1"] == 11
    assert data["2"] == 6
    assert data["3"] == 6
    assert data["4"] == 29

    # Check Info sheet has template name
    assert "Info" in wb.sheetnames
    info_sheet = wb["Info"]
    info_data = {
        info_sheet.cell(r, 1).value: info_sheet.cell(r, 2).value
        for r in range(1, info_sheet.max_row + 1)
    }
    assert "แม่แบบกระดาษคำตอบ" in info_data
    assert "Default #2" in info_data["แม่แบบกระดาษคำตอบ"]


def test_trash_and_permanent_purge(tmp_path):
    app = initialize(tmp_path / "app_data")
    exam = app.exams.create(
        ExamDetails(
            name="ข้อสอบทดสอบลบ",
            academic_year="2569",
            grade="ป.1",
            room="1",
            subject="test",
            question_count=10,
        )
    )
    # Initially in list_exams
    assert any(e.id == exam.id for e in app.exams.list_exams())
    assert not any(e.id == exam.id for e in app.exams.list_archived())

    # Archive (move to trash)
    app.exams.archive(exam.id)
    assert not any(e.id == exam.id for e in app.exams.list_exams())
    assert any(e.id == exam.id for e in app.exams.list_archived())

    # Restore from trash
    app.exams.restore(exam.id)
    assert any(e.id == exam.id for e in app.exams.list_exams())
    assert not any(e.id == exam.id for e in app.exams.list_archived())

    # Move to trash again and permanently purge
    app.exams.archive(exam.id)
    app.exams.purge(exam.id)
    assert not any(e.id == exam.id for e in app.exams.list_exams())
    assert not any(e.id == exam.id for e in app.exams.list_archived())

    # Purging non-existent exam raises ValueError
    with pytest.raises(ValueError):
        app.exams.purge(exam.id)
