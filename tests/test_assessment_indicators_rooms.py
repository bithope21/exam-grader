import sqlite3

import pytest
from openpyxl import load_workbook
from PySide6.QtGui import QImage
from PySide6.QtWidgets import QApplication

from exam_grader.app import initialize
from exam_grader.domain import ExamDetails
from exam_grader.exporting import export_results
from exam_grader.imports import ImportService
from exam_grader.workflow import Workflow


def _image(path, color):
    image = QImage(30, 40, QImage.Format.Format_RGB32)
    image.fill(color)
    assert image.save(str(path))


def _exam_with_records(tmp_path, question_count=4):
    app = initialize(tmp_path / "data")
    exam = app.exams.create(
        ExamDetails("สอบ", "2569", "ม.4", "ป.1/1", "วิชา", question_count)
    )
    importer = ImportService(app.exams.path)
    key_path = tmp_path / "key.png"
    student_path = tmp_path / "student.png"
    _image(key_path, 0xFFFDFFFF)
    _image(student_path, 0xFF000000)
    key_source = importer.import_file(exam.id, key_path, "key")
    student_source = importer.import_file(exam.id, student_path, "student")
    flow = Workflow(app.exams.path)
    key = flow.approve_key(exam.id, ["A"] * question_count, key_source["id"])
    return app, exam, importer, flow, key, student_source


def test_indicators_save_reload_validate_and_score_without_changing_total(tmp_path):
    _app, exam, _importer, flow, key, source = _exam_with_records(tmp_path)
    saved = flow.save_assessment_indicators(
        exam.id,
        [
            {"identifier": "  ว 2.1  ", "from_question": 1, "to_question": 2},
            {"identifier": "การอ่าน", "from_question": 4, "to_question": 4},
        ],
    )
    assert [item["identifier"] for item in saved] == ["ว 2.1", "การอ่าน"]
    assert flow.list_assessment_indicators(exam.id)[0]["from_question"] == 1

    flow.review(source["id"], "1", ["A", "A", "B", "A"], key["id"])
    snapshot = flow.snapshot(exam.id)
    result = snapshot["results"][0]
    assert result["indicator_scores"] == [2, 1]
    assert result["score"] == 3

    with pytest.raises(ValueError, match="ซ้ำ"):
        flow.save_assessment_indicators(
            exam.id,
            [
                {"identifier": "2.1", "from_question": 1, "to_question": 1},
                {"identifier": " 2.1 ", "from_question": 3, "to_question": 3},
            ],
        )
    with pytest.raises(ValueError, match="ซ้อนทับ"):
        flow.save_assessment_indicators(
            exam.id,
            [
                {"identifier": "a", "from_question": 1, "to_question": 2},
                {"identifier": "b", "from_question": 2, "to_question": 3},
            ],
        )
    with pytest.raises(ValueError, match="ระหว่าง 1 ถึง 4"):
        flow.save_assessment_indicators(
            exam.id, [{"identifier": "c", "from_question": 1, "to_question": 5}]
        )


def test_indicator_export_order_and_room_isolation(tmp_path):
    app, exam, importer, flow, key, first_source = _exam_with_records(tmp_path, 4)
    flow.save_assessment_indicators(
        exam.id,
        [
            {"identifier": "2.1", "from_question": 1, "to_question": 2},
            {"identifier": "2.2", "from_question": 3, "to_question": 4},
        ],
    )
    rooms = app.exams.list_rooms(exam.id)
    second_room = app.exams.create_room(exam.id, "ป.1/2")
    second_path = tmp_path / "student-room-2.png"
    _image(second_path, 0xFF00FF00)
    second_source = importer.import_file(
        exam.id, second_path, "student", room_id=second_room.id
    )
    flow.review(first_source["id"], "1", ["A", "B", "A", "B"], key["id"])
    flow.review(second_source["id"], "1", ["B", "B", "A", "A"], key["id"])

    first_snapshot = flow.snapshot(exam.id, rooms[0].id)
    second_snapshot = flow.snapshot(exam.id, second_room.id)
    assert first_snapshot["results"][0]["indicator_scores"] == [1, 1]
    assert second_snapshot["results"][0]["indicator_scores"] == [0, 2]
    assert first_snapshot["results"][0]["score"] == 2
    assert second_snapshot["results"][0]["score"] == 2
    assert len(first_snapshot["results"]) == len(second_snapshot["results"]) == 1

    first_export = export_results(flow, exam.id, output_root=tmp_path / "results", room_id=rooms[0].id)
    workbook = load_workbook(first_export / "scores.xlsx")
    sheet = workbook.active
    assert [sheet.cell(1, column).value for column in range(1, 9)] == [
        "No.",
        "ตัวชี้วัด 2.1",
        "ตัวชี้วัด 2.2",
        "Score",
        "Max",
        "Status",
        "Source File",
        "Review Issues",
    ]
    assert [sheet.cell(2, column).value for column in range(1, 6)] == [1, 1, 1, 2, 4]
    assert [sheet.cell(2, column).data_type for column in range(1, 6)] == ["n"] * 5
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:H2"
    assert sheet[1][0].font.bold is True
    assert sheet[1][0].fill.fgColor.rgb.endswith("EAF2F8")
    assert sheet[1][0].alignment.horizontal == "center"
    assert sheet[1][0].alignment.vertical == "center"
    assert sheet[1][0].border.bottom.style == "thin"
    assert sheet.column_dimensions["A"].width <= 9
    assert sheet.column_dimensions["H"].width <= 42
    assert "/" not in first_export.parent.parent.name
    second_export = export_results(
        flow, exam.id, output_root=tmp_path / "results", room_id=second_room.id
    )
    assert second_export != first_export
    assert second_export.parent.parent.name != first_export.parent.parent.name
    assert len(sqlite3.connect(app.exams.path).execute(
        "SELECT 1 FROM export_runs WHERE exam_id=? AND room_id=?",
        (exam.id, second_room.id),
    ).fetchall()) == 1

    collision_room = app.exams.create_room(exam.id, "ป.1:1")
    collision_path = tmp_path / "student-room-collision.png"
    _image(collision_path, 0xFF0000FF)
    collision_source = importer.import_file(
        exam.id, collision_path, "student", room_id=collision_room.id
    )
    flow.review(collision_source["id"], "1", ["A", "A", "A", "A"], key["id"])
    collision_export = export_results(
        flow, exam.id, output_root=tmp_path / "results", room_id=collision_room.id
    )
    assert collision_export.parent.parent.name not in {
        first_export.parent.parent.name,
        second_export.parent.parent.name,
    }


def test_no_indicators_preserves_legacy_excel_header(tmp_path):
    _app, exam, _importer, flow, key, source = _exam_with_records(tmp_path, 4)
    flow.review(source["id"], "1", ["A", "B", "A", "B"], key["id"])
    path = export_results(flow, exam.id, output_root=tmp_path / "results")
    workbook = load_workbook(path / "scores.xlsx")
    assert [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))] == [
        "No.",
        "Score",
        "Max",
        "Status",
        "Source File",
        "Review Issues",
    ]


def test_invalid_indicator_after_question_count_change_does_not_break_core_snapshot(tmp_path):
    _app, exam, _importer, flow, key, source = _exam_with_records(tmp_path, 4)
    flow.save_assessment_indicators(
        exam.id, [{"identifier": "2.1", "from_question": 3, "to_question": 4}]
    )
    flow.review(source["id"], "1", ["A", "A", "A", "A"], key["id"])
    # This is intentionally an invalid persisted configuration; the migration-safe
    # behavior is to preserve it, keep total grading available, and block indicator export.
    _app.exams.update_exam_template(exam.id, "default-1", 1, question_count=2)
    snapshot = flow.snapshot(exam.id)
    assert snapshot["assessment_indicators"]["valid"] is False
    assert snapshot["results"][0]["score"] == 4
    with pytest.raises(ValueError, match="ตัวชี้วัด"):
        export_results(flow, exam.id, output_root=tmp_path / "results")


def test_v14_exam_room_data_migrates_without_loss(tmp_path):
    database = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.executescript(
            """
            CREATE TABLE exams (
                id TEXT PRIMARY KEY, name TEXT NOT NULL, academic_year TEXT NOT NULL,
                grade TEXT NOT NULL, room TEXT NOT NULL, subject TEXT NOT NULL,
                question_count INTEGER NOT NULL, expected_number_max INTEGER,
                archived_at TEXT, output_root TEXT, template_id TEXT NOT NULL,
                template_version INTEGER NOT NULL, created_at TEXT NOT NULL
            );
            CREATE TABLE sources (
                id TEXT PRIMARY KEY, exam_id TEXT NOT NULL, sha256 TEXT NOT NULL,
                original_name TEXT NOT NULL, relative_path TEXT NOT NULL,
                width INTEGER NOT NULL, height INTEGER NOT NULL, created_at TEXT NOT NULL,
                purpose TEXT NOT NULL, archived_at TEXT, UNIQUE(exam_id, sha256)
            );
            CREATE TABLE import_failures (
                id TEXT PRIMARY KEY, exam_id TEXT NOT NULL, path TEXT NOT NULL,
                purpose TEXT NOT NULL, error TEXT NOT NULL, created_at TEXT NOT NULL,
                resolved_at TEXT
            );
            CREATE TABLE export_runs (
                id TEXT PRIMARY KEY, exam_id TEXT NOT NULL, path TEXT NOT NULL,
                run_id TEXT NOT NULL, snapshot_fingerprint TEXT NOT NULL,
                created_at TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'published'
            );
            CREATE TABLE attendance (
                exam_id TEXT NOT NULL, student_number INTEGER NOT NULL,
                status TEXT NOT NULL, updated_at TEXT NOT NULL,
                PRIMARY KEY(exam_id, student_number)
            );
            CREATE TABLE skipped_numbers (
                exam_id TEXT NOT NULL, student_number INTEGER NOT NULL,
                PRIMARY KEY(exam_id, student_number)
            );
            INSERT INTO exams VALUES
                ('exam-1','สอบ','2569','ป.1','ป.1/1','คณิต',4,NULL,NULL,NULL,'default-1',1,'2026');
            INSERT INTO sources VALUES
                ('key-1','exam-1','key-hash','key.png','input/key',30,40,'2026','key',NULL),
                ('student-1','exam-1','student-hash','student.png','input/student',30,40,'2026','student',NULL);
            INSERT INTO import_failures VALUES
                ('failure-1','exam-1','student.png','student','old error','2026',NULL);
            INSERT INTO export_runs VALUES
                ('export-1','exam-1','results/old','run-1','fingerprint','2026','published');
            INSERT INTO attendance VALUES ('exam-1',7,'absent','2026');
            INSERT INTO skipped_numbers VALUES ('exam-1',8);
            PRAGMA user_version = 14;
            """
        )

    from exam_grader.storage import ExamStore

    store = ExamStore(database)
    store.initialize()
    room = store.list_rooms("exam-1")[0]
    with sqlite3.connect(database) as connection:
        connection.row_factory = sqlite3.Row
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 15
        assert connection.execute(
            "SELECT room_id FROM sources WHERE id='student-1'"
        ).fetchone()[0] == room.id
        assert connection.execute(
            "SELECT room_id FROM attendance WHERE student_number=7"
        ).fetchone()[0] == room.id
        assert connection.execute(
            "SELECT room_id FROM skipped_numbers WHERE student_number=8"
        ).fetchone()[0] == room.id
        assert connection.execute(
            "SELECT room_id FROM export_runs WHERE id='export-1'"
        ).fetchone()[0] == room.id


def test_indicator_dialog_and_optional_room_creation(tmp_path):
    qt = QApplication.instance() or QApplication([])
    app = initialize(tmp_path / "data")
    exam = app.exams.create(ExamDetails("สอบ", "2569", "ป.1", "", "คณิต", 4))
    assert exam.details.room == "ห้อง 1"
    assert app.exams.list_rooms(exam.id)[0].label == "ห้อง 1"

    from exam_grader.exam_ui import AssessmentIndicatorsDialog
    from exam_grader.ui import NewExamDialog

    dialog = AssessmentIndicatorsDialog(app.exams.path, exam.id, 4)
    dialog._add_row()
    dialog._add_row()
    dialog.rows[0][1].setText("2.1")
    dialog.rows[0][2].setValue(1)
    dialog.rows[0][3].setValue(2)
    dialog.rows[1][1].setText("การอ่าน")
    dialog.rows[1][2].setValue(4)
    dialog.rows[1][3].setValue(4)
    dialog.accept()
    assert [item["identifier"] for item in Workflow(app.exams.path).list_assessment_indicators(exam.id)] == [
        "2.1",
        "การอ่าน",
    ]
    dialog.close()

    new_dialog = NewExamDialog(application=app)
    assert new_dialog.fields["room"].currentText() == ""
    new_dialog.close()
    qt.processEvents()
