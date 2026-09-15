import json

import cv2
import numpy as np
from openpyxl import load_workbook
from PySide6.QtGui import QImage

from exam_grader.app import initialize
from exam_grader.domain import ExamDetails
from exam_grader.exporting import _partial_status_banner, checked_filename, export_results
from exam_grader.imaging import OMR_PIPELINE_VERSION
from exam_grader.imports import ImportService
from exam_grader.review_service import ReviewService
from exam_grader.workflow import Workflow


def _write_image(path, marker):
    image = QImage(80, 100, QImage.Format.Format_RGB32)
    image.fill(0xFFF5F5F5)
    image.setPixel(0, 0, 0xFF000000 | (marker << 16))
    image.save(str(path))


def test_partial_review_exports_unknown_identity_without_guessing(tmp_path):
    app = initialize(tmp_path / "data")
    exam = app.exams.create(ExamDetails("สอบบางส่วน", "2569", "ม.4", "1", "วิชา", 3))
    importer = ImportService(app.exams.path)
    key_path, first_path, second_path = (
        tmp_path / "key.png",
        tmp_path / "first.png",
        tmp_path / "second.png",
    )
    for marker, path in enumerate((key_path, first_path, second_path), start=1):
        _write_image(path, marker)
    key_source = importer.import_file(exam.id, key_path, "key")
    first = importer.import_file(exam.id, first_path, "student")
    second = importer.import_file(exam.id, second_path, "student")
    flow = Workflow(app.exams.path)
    flow.approve_key(exam.id, ["A", "B", "C"], key_source["id"])
    identity = [[1, 0, 0], [0, 1, 0], [0, 0, 1]]
    flow.save_detection(
        first["id"],
        {
            "pipeline_version": OMR_PIPELINE_VERSION,
            "registration": {"matrix": identity},
            "student_number_observation": {"candidate": "9", "candidates": ["9"]},
            "answers": [
                {"classification": "single_mark", "selected": ["A"], "auto_resolved": True},
                {"classification": "uncertain", "selected": ["B"], "auto_resolved": False},
                {"classification": "multiple", "selected": ["A", "C"], "auto_resolved": True},
            ],
        },
    )
    flow.save_detection(
        second["id"],
        {
            "pipeline_version": OMR_PIPELINE_VERSION,
            "registration": {"matrix": identity},
            "student_number_observation": {"candidate": "8", "candidates": ["8"]},
            "answers": [
                {"classification": "single_mark", "selected": ["A"], "auto_resolved": True},
                {"classification": "blank", "selected": [], "auto_resolved": True},
                {"classification": "single_mark", "selected": ["C"], "auto_resolved": True},
            ],
        },
    )
    service = ReviewService(app.exams.path)
    answer_edits = {first["id"]: {2: "B"}}
    identity_edits = {first["id"]: "1"}
    assert service.skip_summary(
        exam.id, answer_edits=answer_edits, identity_edits=identity_edits
    ) == {
        "unconfirmed_identities": 1,
        "unclear_answers": 1,
        "unresolved_sheets": 2,
    }

    result = service.skip_remaining(
        exam.id, answer_edits=answer_edits, identity_edits=identity_edits
    )
    assert result == {"partial_sheets": 2, "unknown_identities": 1}
    snapshot = flow.snapshot(exam.id)
    first_result, unknown_result = snapshot["results"]
    assert first_result["student_number"] == "1"
    assert first_result["status"] == "review_skipped"
    assert first_result["answers"] == ["A", "B", "unresolved"]
    assert first_result["score"] == 2
    assert first_result["answer_provenance"][1] == {"question": 2, "origin": "teacher_confirmed"}
    assert "Q3 unresolved" in first_result["review_issues"]
    assert unknown_result["student_number"] == "ไม่ทราบเลขที่ #1"
    assert unknown_result["identity_confirmed"] is False
    assert unknown_result["score"] == 2

    # Re-skipping and re-exporting retain the same database label and checked filename.
    service.skip_remaining(exam.id)
    second_snapshot = flow.snapshot(exam.id)
    assert second_snapshot["results"][1]["student_number"] == "ไม่ทราบเลขที่ #1"
    first_export = export_results(flow, exam.id, tmp_path / "out")
    second_export = export_results(flow, exam.id, tmp_path / "out")
    assert (first_export / "checked" / "ไม่ทราบเลขที่-01.jpg").exists()
    first_json = json.loads((first_export / "_system" / "results.json").read_text())
    assert first_json["results"][1]["checked_image"] == "checked/ไม่ทราบเลขที่-01.jpg"
    assert first_json["results"][1]["checked_provenance"]["identity_confirmed"] is False
    assert "identity_unknown" in first_json["results"][1]["review_issues"]
    assert (second_export / "checked" / "ไม่ทราบเลขที่-01.jpg").exists()

    workbook = load_workbook(first_export / "scores.xlsx", data_only=True)
    sheet = workbook["Scores"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[-2:] == ["Source File", "Review Issues"]
    assert sheet.cell(2, 4).value == "review_skipped"
    assert "Q3 unresolved" in sheet.cell(2, 6).value
    checked = cv2.imread(str(first_export / "checked" / "ไม่ทราบเลขที่-01.jpg"))
    assert checked is not None and checked.shape[0] > 100


def test_unknown_checked_filenames_are_stable_and_collision_safe():
    used = set()
    assert checked_filename("ไม่ทราบเลขที่ #1", used) == "ไม่ทราบเลขที่-01.jpg"
    assert checked_filename("ไม่ทราบเลขที่ #2", used) == "ไม่ทราบเลขที่-02.jpg"
    assert checked_filename("#1", used) == "ไม่ทราบเลขที่-03.jpg"


def test_partial_checked_banner_renders_thai_text():
    image = np.full((60, 120, 3), 255, np.uint8)
    result = {
        "status": "review_skipped",
        "review_issues": ["review_skipped", "identity_unknown", "Q8 unresolved"],
    }
    bannered = _partial_status_banner(image, result)
    assert bannered.shape == (102, 120, 3)
    assert np.count_nonzero(np.any(bannered[:42] < 245, axis=2)) > 100
