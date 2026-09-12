"""Focused Windows UAT verification tests for v1.0.1 hardening.

Tests:
1. Settings -> Template Management loads built-ins and reference images without crash.
2. Grade -> Save results + Excel exports cleanly to real Documents path including Thai/Unicode without WinError 5.
3. Verify output integrity: checked results complete, Excel opens with openpyxl, no .staging-* directories remaining, source originals untouched.
"""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import cv2
import openpyxl
import pytest
from PySide6.QtWidgets import QApplication

from exam_grader.app import initialize
from exam_grader.domain import ExamDetails
from exam_grader.exporting import export_results
from exam_grader.imports import ImportService
from exam_grader.settings_ui import TemplateSettingsDialog
from exam_grader.template_manager import (
    BUILTIN_TEMPLATE_IDS,
    get_reference_image,
    load_builtin_template,
)
from exam_grader.workflow import Workflow


@pytest.fixture(autouse=True)
def ensure_qapp():
    return QApplication.instance() or QApplication([])


def test_template_management_dialog_and_resources(tmp_path: Path):
    """Focused UAT: Template Management opens, loads all built-ins and reference images."""
    app = initialize(tmp_path / "app_data")

    # 1. Verify all 3 built-in definitions load
    for b_id in BUILTIN_TEMPLATE_IDS:
        t = load_builtin_template(b_id)
        assert t.template_id == b_id
        assert t.kind == "builtin"
        assert t.question_count in (30, 60)

        # Reference image must load and verify sha256
        ref = get_reference_image(t, app.data_dir)
        assert ref is not None
        assert ref.size > 0

    # 2. Verify dialog initializes without crash and lists all built-in templates
    dlg = TemplateSettingsDialog(app)
    assert dlg.table.rowCount() >= 3

    names = [dlg.table.item(r, 0).text() for r in range(dlg.table.rowCount()) if dlg.table.item(r, 0)]
    assert any("Default #1" in n for n in names)
    assert any("Default #2" in n for n in names)
    assert any("Default #3" in n for n in names)

    # 3. Simulate selection of each row
    for r in range(dlg.table.rowCount()):
        dlg.table.selectRow(r)
        dlg._on_selection_changed()
        assert dlg.selected_template is not None
        assert dlg.preview_image_lbl.pixmap() is not None

    dlg.close()


def test_windows_export_real_documents_thai_unicode_path():
    """Focused UAT: Grade -> Save results + Excel export to real Documents Thai path without WinError 5."""
    docs_dir = Path.home() / "Documents"
    if not docs_dir.exists():
        docs_dir = Path("D:/AI/workspaces/test_documents")
        docs_dir.mkdir(parents=True, exist_ok=True)

    # Real Thai/Unicode folder under Documents
    test_export_root = docs_dir / "ทดสอบตรวจข้อสอบ_ไทย_๒๕๖๙"
    test_export_root.mkdir(parents=True, exist_ok=True)

    # Create isolated app data for clean test
    app_data = test_export_root / ".test_app_data"
    if app_data.exists():
        shutil.rmtree(app_data)

    app = initialize(app_data)

    try:
        exam = app.exams.create(
            ExamDetails(
                name="สอบกลางภาค ภาษาไทย",
                academic_year="2569",
                subject="ภาษาไทยเพื่อการสื่อสาร",
                grade="ม.๑",
                room="๓",
                question_count=30,
                template_id="default-3",
                template_version=1,
            )
        )

        importer = ImportService(app.exams.path)
        flow = Workflow(app.exams.path)
        t3 = load_builtin_template("default-3")

        # Use real fixture image from vol.6
        vol6_dir = Path("tests/fixtures/real/vol.6")
        if not (vol6_dir / "key.JPG").exists():
            pytest.skip("vol.6 fixtures not present")

        # Compute SHA-256 of source images before import
        key_source_bytes = (vol6_dir / "key.JPG").read_bytes()
        key_sha_before = hashlib.sha256(key_source_bytes).hexdigest()

        student_files = sorted(vol6_dir.glob("IMG_*.JPG"))[:3]
        student_shas_before = {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in student_files}

        # 1. Import key
        k_src = importer.import_file(exam.id, vol6_dir / "key.JPG", "key")
        k_ans = [
            "B", "B", "B", "A", "D", "C", "B", "B", "C", "B",
            "A", "D", "C", "D", "A", "B", "B", "D", "B", "C",
            "A", "B", "C", "B", "A", "D", "A", "B", "D", "B",
        ]
        approved_key = flow.approve_key(exam.id, k_ans, k_src["id"])

        # 2. Import student sheets & simulate detection + review
        for idx, sf in enumerate(student_files, start=1):
            s_src = importer.import_file(exam.id, sf, "student")
            flow.save_detection(
                s_src["id"],
                {
                    "registration": {
                        "matrix": [[1.0, 0.0, 0.0], [0.0, 1.0, 0.0], [0.0, 0.0, 1.0]],
                        "table_coverage": 1.0,
                    },
                    "answers": [{"selected": ["A"], "confidence": 0.95} for _ in range(30)],
                },
            )
            flow.review(
                s_src["id"],
                str(idx),
                ["A"] * 30,
                approved_key["id"],
            )

        # 3. Perform export to Thai/Unicode path in Documents
        exported_path = export_results(flow, exam.id, output_root=test_export_root)

        # Verification 1: Export path exists
        assert exported_path.exists()
        assert exported_path.is_dir()

        # Verification 2: Excel file exists and is valid
        excel_files = list(exported_path.glob("*.xlsx"))
        assert len(excel_files) == 1, f"Expected 1 excel file, found {len(excel_files)}"
        excel_path = excel_files[0]
        wb = openpyxl.load_workbook(excel_path)
        assert len(wb.sheetnames) >= 1
        ws = wb.active
        assert ws is not None
        # Check title row or content
        wb.close()

        # Verification 3: Checked images exist for each student in checked/ folder
        checked_images = list((exported_path / "checked").glob("*.jpg"))
        assert len(checked_images) == len(student_files), (
            f"Expected {len(student_files)} checked images in checked/, found {len(checked_images)}"
        )

        # Verification 4: No .staging-* folders remain anywhere in export root
        staging_folders = list(test_export_root.rglob(".staging-*"))
        assert len(staging_folders) == 0, f"Found leftover staging folders: {staging_folders}"

        # Verification 5: Source files are completely unmodified (byte-for-byte SHA256 match)
        assert hashlib.sha256((vol6_dir / "key.JPG").read_bytes()).hexdigest() == key_sha_before
        for f in student_files:
            assert hashlib.sha256(f.read_bytes()).hexdigest() == student_shas_before[f.name]

    finally:
        # Clean up test output directory
        if test_export_root.exists():
            shutil.rmtree(test_export_root, ignore_errors=True)


def test_windows_uat_workspace_fixture_and_excel_integrity():
    """Verify full UAT using D:\\AI\\workspaces\\exam-grader-test images and Thai path in Documents."""
    ws_dir = Path("D:/AI/workspaces/exam-grader-test")
    if not (ws_dir / "IMG_0910.JPG").exists():
        pytest.skip("Workspace test images not available")

    docs_dir = Path.home() / "Documents"
    uat_out_dir = docs_dir / "ผลการตรวจ_UAT_ป1"
    uat_out_dir.mkdir(parents=True, exist_ok=True)

    app_data = uat_out_dir / ".app_data"
    if app_data.exists():
        shutil.rmtree(app_data)
    app = initialize(app_data)

    try:
        exam = app.exams.create(
            ExamDetails(
                name="วิชาศิลปะและดนตรี",
                academic_year="2569",
                subject="ศิลปะ",
                grade="ป.๑",
                room="๑",
                question_count=30,
                template_id="default-3",
                template_version=1,
            )
        )
        importer = ImportService(app.exams.path)
        flow = Workflow(app.exams.path)

        # Hash source images to verify immutability
        source_files = [
            ws_dir / "IMG_0910.JPG",
            ws_dir / "IMG_0911.JPG",
            ws_dir / "IMG_0912.JPG",
            ws_dir / "IMG_0913.JPG",
            ws_dir / "IMG_0914.JPG",
        ]
        shas_before = {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in source_files}

        # Import teacher key (IMG_0910.JPG)
        k_src = importer.import_file(exam.id, ws_dir / "IMG_0910.JPG", "key")
        key_ans = [
            "B", "B", "B", "A", "D", "C", "B", "B", "C", "B",
            "A", "D", "C", "D", "A", "B", "B", "D", "B", "C",
            "A", "B", "C", "B", "A", "D", "A", "B", "D", "B",
        ]
        approved_key = flow.approve_key(exam.id, key_ans, k_src["id"])

        # Import 4 students (IMG_0911 - IMG_0914)
        student_files = [ws_dir / f"IMG_091{i}.JPG" for i in range(1, 5)]
        for s_idx, s_file in enumerate(student_files, start=1):
            s_src = importer.import_file(exam.id, s_file, "student")
            flow.save_detection(
                s_src["id"],
                {
                    "registration": {
                        "matrix": [[1.0, 0.0, 0.0], [0.0, 1.0, 0.0], [0.0, 0.0, 1.0]],
                        "table_coverage": 1.0,
                    },
                    "answers": [{"selected": ["B"], "confidence": 0.99} for _ in range(30)],
                },
            )
            flow.review(
                s_src["id"],
                str(s_idx),
                ["B"] * 30,
                approved_key["id"],
            )

        # Export to Documents Thai path
        final_export = export_results(flow, exam.id, output_root=uat_out_dir)

        # 1. Directory exists
        assert final_export.exists() and final_export.is_dir()

        # 2. No staging directories left anywhere
        staging_dirs = list(uat_out_dir.rglob(".staging-*"))
        assert len(staging_dirs) == 0, f"Found staging dirs: {staging_dirs}"

        # 3. 4 checked images exist in checked/ folder
        checked_imgs = list((final_export / "checked").glob("*.jpg"))
        assert len(checked_imgs) == 4, f"Expected 4 checked images, got {len(checked_imgs)}"

        # 4. Excel opens and has correct sheets and rows
        excel_path = next(final_export.glob("*.xlsx"))
        wb = openpyxl.load_workbook(excel_path)
        assert "Scores" in wb.sheetnames
        assert "Info" in wb.sheetnames
        scores_ws = wb["Scores"]
        # Header + 4 student rows
        assert scores_ws.max_row == 5
        assert scores_ws.cell(1, 1).value == "No."
        assert scores_ws.cell(1, 2).value == "Score"
        wb.close()

        # 5. Source originals are 100% untouched
        for f in source_files:
            assert hashlib.sha256(f.read_bytes()).hexdigest() == shas_before[f.name], (
                f"Source file {f.name} was modified!"
            )

    finally:
        if uat_out_dir.exists():
            shutil.rmtree(uat_out_dir, ignore_errors=True)

