"""Replay existing teacher decisions in a fresh disposable real-fixture UI run.

This is engineering UAT, never a new teacher confirmation or blind accuracy trial.
Run with QT_QPA_PLATFORM=offscreen for exact logical viewport sizes, or cocoa for
the native desktop including actual maximization. All images are derived copies.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
from pathlib import Path

import cv2
import numpy as np
from openpyxl import load_workbook
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QApplication, QDialog, QLineEdit

from exam_grader.app import initialize
from exam_grader.domain import ExamDetails
from exam_grader.exam_ui import BatchWorker, ExamDialog
from exam_grader.exporting import export_results
from exam_grader.identity import observe
from exam_grader.imaging import decode, template
from exam_grader.imports import ImportService
from exam_grader.review_ui import ReviewDialog, key_answer_text
from exam_grader.workflow import Workflow


def settle():
    loop = QEventLoop()
    QTimer.singleShot(150, loop.quit)
    loop.exec()


def capture(dialog: ReviewDialog, output: Path, prefix: str) -> list[dict]:
    records = []
    for width, height in ((1024, 768), (1440, 900)):
        dialog.showNormal()
        dialog.resize(width, height)
        dialog.show()
        settle()
        path = output / f"{prefix}-{width}x{height}.png"
        assert dialog.grab().save(str(path))
        record = {"requested": [width, height], "actual": [dialog.width(), dialog.height()],
                  "dpr": dialog.devicePixelRatioF(), "screenshot": str(path),
                  "table_scroll_max": dialog.table.horizontalScrollBar().maximum(),
                  "column_widths": [dialog.table.columnWidth(i) for i in range(3)],
                  "table_viewport_width": dialog.table.viewport().width()}
        assert record["table_scroll_max"] == 0
        assert sum(record["column_widths"]) <= record["table_viewport_width"]
        records.append(record)
    if QApplication.platformName() != "offscreen":
        dialog.showMaximized()
        settle()
        path = output / f"{prefix}-maximized.png"
        assert dialog.grab().save(str(path))
        records.append({"maximized": dialog.isMaximized(), "actual": [dialog.width(), dialog.height()],
                        "dpr": dialog.devicePixelRatioF(), "screenshot": str(path)})
    return records


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--truth", type=Path, default=Path("docs/evidence/polish/truth.json"))
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    truth = json.loads(args.truth.read_text())
    fixture_root = Path("tests/fixtures/real")
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in fixture_root.iterdir() if p.is_file()}
    qt = QApplication.instance() or QApplication([])
    qt.setApplicationName("Exam Grader · Engineering UAT")
    app = initialize(Path(tempfile.mkdtemp(prefix="exam-grader-polish-")))
    exam = app.exams.create(ExamDetails("ทดสอบกระดาษจริง", "2569", "ม.6", "1", "คณิตศาสตร์", 30))
    flow = Workflow(app.exams.path)
    importer = ImportService(app.exams.path)
    app.exams.set_output_root(exam.id, output / "exports")
    reports: dict = {"kind": "engineering UI replay of historical teacher decisions",
                     "platform": QApplication.platformName(), "truth_file": str(args.truth.resolve()),
                     "data_dir": str(app.data_dir), "exam_id": exam.id, "ui": [], "numbers": []}
    for purpose, paths in (("key", [fixture_root / "key.JPG"]),
                           ("student", sorted(fixture_root.glob("IMG*")))):
        worker = BatchWorker(app.exams.path, exam.id, paths, purpose)
        worker.start()
        assert worker.wait(90000)
        assert not importer.list_failures(exam.id)
        if purpose == "key":
            source = importer.list_sources(exam.id)[0]
            dialog = ReviewDialog(app.exams.path, source)
            for combo, answer in zip(dialog.combos, truth["key"]["answers"]):
                if isinstance(combo, QLineEdit):
                    combo.setText(key_answer_text(answer))
                else:
                    combo.setCurrentIndex(combo.findData(answer))
            reports["ui"] += capture(dialog, output, "key")
            dialog.confirmed.setChecked(True)
            dialog.accept()
            assert dialog.result() == QDialog.DialogCode.Accepted
    by_name = {r["filename"]: r for r in truth["records"]}
    for source in importer.list_sources(exam.id):
        if source["purpose"] != "student":
            continue
        row = by_name[source["original_name"]]
        assert source["sha256"] == row["source_sha256"]
        detection = flow.latest_detection(source["id"])
        observed = observe(importer.verified_bytes(source), detection["registration"]["matrix"],
                           diagnostics_dir=output / "numbers" / Path(row["filename"]).stem)
        assert observed == detection["student_number_observation"]
        reports["numbers"].append({"filename": row["filename"], "truth": row["student_number"],
                                  "correct": observed["candidate"] == row["student_number"],
                                  "truth_review_id": row["review_id"], "observation": observed})
        dialog = ReviewDialog(app.exams.path, source)
        dialog.number.setText(row["student_number"])
        for combo, answer in zip(dialog.combos, row["answers"]):
            combo.setCurrentIndex(combo.findData(answer))
        if row["student_number"] == "2":
            reports["ui"] += capture(dialog, output, "student")
        dialog.confirmed.setChecked(True)
        dialog.accept()
        assert dialog.result() == QDialog.DialogCode.Accepted
    final = export_results(flow, exam.id, app.exams.output_root(exam.id))
    reports["output"] = str(final)
    snapshot = json.loads((final / "_system/results.json").read_text())
    workbook = load_workbook(final / "scores.xlsx")
    reports["scores"] = [r["score"] for r in snapshot["results"]]
    assert reports["scores"] == [10, 6, 12, 26, 8]
    assert [workbook.active.cell(i + 2, 2).value for i in range(5)] == reports["scores"]
    manifest = json.loads((final / "_system/manifest.json").read_text())
    for name, digest in manifest.items():
        assert hashlib.sha256((final / name).read_bytes()).hexdigest() == digest
    reports["manifest_hashes_verified"] = len(manifest)
    reports["directory_tree"] = sorted(str(p.relative_to(final)) for p in final.rglob("*"))
    headers, grids = [], []
    for result in snapshot["results"]:
        image = decode((final / result["checked_image"]).read_bytes())
        aligned = cv2.warpPerspective(image, np.array(result["detection"]["registration"]["matrix"]),
                                      (template()["width"], template()["height"]))
        header = cv2.resize(aligned[:145, 520:], None, fx=3, fy=3)
        grid = cv2.resize(aligned[227:714, 18:425], None, fx=1.5, fy=1.5)
        stem = Path(result["source"]["original_name"]).stem
        cv2.imwrite(str(output / f"{stem}-checked-header.png"), header)
        cv2.imwrite(str(output / f"{stem}-checked-grid.png"), grid)
        headers.append(header)
        grids.append(grid)
    cv2.imwrite(str(output / "checked-headers.png"), np.vstack(headers))
    # Reopen from storage and verify the real reveal control targets the visible
    # run directory. Opening Finder is separately exercised during native UAT.
    reopened = ExamDialog(initialize(app.data_dir), exam)
    reopened.tabs.setCurrentIndex(3)
    reopened.show()
    settle()
    assert reopened.export_history.currentData() == str(final)
    assert reopened.reveal_button.isEnabled()
    reopened.grab().save(str(output / "results.png"))
    if QApplication.platformName() != "offscreen":
        reopened.reveal_result()
        settle()
    reopened.close()
    assert hashes == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in fixture_root.iterdir() if p.is_file()}
    reports["original_hashes_unchanged"] = True
    reports["review_required"] = 5
    reports["auto_accept"] = 0
    reports["candidate_correct"] = sum(r["correct"] for r in reports["numbers"])
    (output / "report.json").write_text(json.dumps(reports, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"output": str(final), "candidate_correct": reports["candidate_correct"],
                      "scores": reports["scores"], "platform": QApplication.platformName()}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
