"""Publish complete, versioned evidence packages from one immutable snapshot."""

import hashlib
import json
import os
import re
import shutil
import tempfile
import unicodedata
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from PySide6.QtGui import QColor, QFont, QImage, QPainter
from PySide6.QtWidgets import QApplication

from exam_grader import __version__
from exam_grader.imaging import cell_rect, decode, template
from exam_grader.imports import ImportService
from exam_grader.preferences import annotation_colors
from exam_grader.storage import ExamStore
from exam_grader.template_manager import TemplateDefinition, load_exam_template_def
from exam_grader.workflow import Workflow, score_answer


# OpenCV uses BGR(A). Keep the legend explicit and versioned with each
# derived checked image so the colours are evidence presentation only; they
# never influence recognition or scoring.
def checked_colors() -> dict[str, tuple[int, int, int, int]]:
    return annotation_colors()


CHECKED_COLORS = checked_colors()
_RENDER_APPLICATION: QApplication | None = None


def _ensure_render_application() -> None:
    global _RENDER_APPLICATION
    if QApplication.instance() is not None:
        return
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    _RENDER_APPLICATION = QApplication([])


def human_readable_name(exam: dict, room_label: str | None = None) -> str:
    room = exam.get("room", "") if room_label is None else room_label
    raw = "_".join(
        str(room if field == "room" else exam.get(field, "")).strip()
        for field in ("academic_year", "grade", "room", "subject", "name")
    )
    normalized = unicodedata.normalize("NFKC", raw)
    safe = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "-", normalized)
    safe = re.sub(r"\s+", " ", safe).strip(" .")
    # Filesystems commonly limit each component to 255 UTF-8 bytes; Thai is
    # three bytes per code point. Leave space for suffixes, not just characters.
    safe = safe.encode("utf-8")[:180].decode("utf-8", errors="ignore").rstrip(" .")
    return (safe or "exam") + f"_{int(exam.get('question_count', 0))}q"


def checked_filename(number: str | None, used: set[str]) -> str:
    normalized = (number or "").translate(str.maketrans("๐๑๒๓๔๕๖๗๘๙", "0123456789"))
    known = (
        normalized.isascii()
        and normalized.isdigit()
        and len(normalized) <= 6
        and int(normalized) > 0
    )
    if known:
        name = f"เลขที่-{int(normalized):02d}.jpg"
        stem = f"เลขที่-{int(normalized):02d}"
        sequence = 2
        while name in used:
            name = f"{stem}_{sequence:02d}.jpg"
            sequence += 1
    else:
        match = re.search(r"#(\d+)$", number or "")
        sequence = int(match.group(1)) if match else 1
        name = f"ไม่ทราบเลขที่-{sequence:02d}.jpg"
        while name in used:
            sequence += 1
            name = f"ไม่ทราบเลขที่-{sequence:02d}.jpg"
    used.add(name)
    return name


def _partial_status_banner(image: np.ndarray, result: dict) -> np.ndarray:
    if result.get("status") != "review_skipped":
        return image
    issues = [issue for issue in result.get("review_issues", []) if issue != "review_skipped"]
    count = max(1, len(issues))
    _ensure_render_application()
    banner = QImage(image.shape[1], 42, QImage.Format.Format_RGB32)
    banner.fill(QColor("white"))
    painter = QPainter(banner)
    label = f"ตรวจทานไม่ครบ · ข้าม {count} รายการ"
    font = QFont()
    font.setPixelSize(22)
    font.setBold(True)
    painter.setFont(font)
    painter.setPen(QColor(0, 110, 220))
    painter.drawText(14, 29, label)
    painter.end()
    banner = banner.convertToFormat(QImage.Format.Format_RGB888)
    pixels = (
        np.frombuffer(banner.bits(), np.uint8)
        .reshape(banner.height(), banner.bytesPerLine())[:, : banner.width() * 3]
        .reshape(banner.height(), banner.width(), 3)
    )
    banner_bgr = cv2.cvtColor(pixels, cv2.COLOR_RGB2BGR)
    return np.vstack((banner_bgr, image))


def checked_image(
    data: bytes,
    result: dict,
    key: dict,
    *,
    template_def: TemplateDefinition | None = None,
) -> bytes:
    original = decode(data)
    detection = result.get("detection") or {}
    registration = detection.get("registration")
    colors = checked_colors()
    if registration and registration.get("matrix"):
        if template_def is not None:
            can_w = template_def.canonical_width
            can_h = template_def.canonical_height
            score_box = template_def.score_roi
        else:
            geometry = template()
            can_w = geometry["width"]
            can_h = geometry["height"]
            score_box = tuple(geometry["score_roi"])

        canonical = np.zeros((can_h, can_w, 4), np.uint8)
        for index, (answer, correct) in enumerate(
            zip(result["answers"], key["answers"], strict=True), start=1
        ):
            if index > result["max"]:
                break
            choice = answer if answer in "ABCDE" else None
            if choice is not None:
                choice_index = "ABCDE".index(choice)
                x, y, w, h = cell_rect(index, choice_index, template_def=template_def)
                color = colors["correct"] if score_answer(answer, correct) else colors["incorrect"]
                if score_answer(answer, correct):
                    cv2.line(canonical, (x + 4, y + h // 2), (x + w // 3, y + h - 5), color, 2)
                    cv2.line(canonical, (x + w // 3, y + h - 5), (x + w - 3, y + 4), color, 2)
                else:
                    cv2.line(canonical, (x + 4, y + 4), (x + w - 4, y + h - 4), color, 2)
                    cv2.line(canonical, (x + w - 4, y + 4), (x + 4, y + h - 4), color, 2)
            elif answer in {"blank", "multiple", "boundary_cross"}:
                x, y, w, h = cell_rect(index, 0, template_def=template_def)
                color = colors.get(answer, colors.get("other", (0, 165, 255, 255)))
                if answer == "boundary_cross":
                    observation = (
                        (detection.get("answers") or [])[index - 1]
                        if index <= len(detection.get("answers") or [])
                        else {}
                    )
                    selected = observation.get("selected") or []
                    if len(selected) >= 2 and all(value in "ABCDE" for value in selected):
                        first = cell_rect(
                            index, "ABCDE".index(selected[0]), template_def=template_def
                        )
                        last = cell_rect(
                            index, "ABCDE".index(selected[-1]), template_def=template_def
                        )
                        x, y = first[0], min(first[1], last[1])
                        w = last[0] + last[2] - x
                        h = max(first[1] + first[3], last[1] + last[3]) - y
                    cv2.rectangle(canonical, (x + 2, y + 2), (x + w - 2, y + h - 2), color, 2)
                    cv2.line(canonical, (x + 4, y + 4), (x + w - 4, y + h - 4), color, 2)
                    cv2.line(canonical, (x + w - 4, y + 4), (x + 4, y + h - 4), color, 2)
                else:
                    cv2.rectangle(canonical, (x + 4, y + 4), (x + w - 4, y + h - 4), color, 2)
                if answer == "multiple":
                    cv2.putText(
                        canonical,
                        "!",
                        (x + w // 2 - 4, y + h - 5),
                        cv2.FONT_HERSHEY_SIMPLEX,
                        0.5,
                        color,
                        2,
                    )
        if score_box is not None:
            sx1, sy1, sx2, sy2 = score_box
            label = f"{result['score']} / {result['max']}"
            font = cv2.FONT_HERSHEY_SIMPLEX
            (unit_width, unit_height), baseline = cv2.getTextSize(label, font, 1, 2)
            scale = min((sx2 - sx1 - 12) / unit_width, (sy2 - sy1 - 12) / (unit_height + baseline))
            (width, height), baseline = cv2.getTextSize(label, font, scale, 2)
            origin = (sx1 + (sx2 - sx1 - width) // 2, sy1 + (sy2 - sy1 + height) // 2)
            cv2.putText(canonical, label, origin, font, scale, colors["score"], 2, cv2.LINE_AA)
        matrix = np.asarray(registration["matrix"], dtype=np.float64)
        inverse = np.linalg.inv(matrix)
        max_dimension = 2400
        longest = max(original.shape[:2])
        output_scale = min(1.0, max_dimension / longest)
        if output_scale < 1.0:
            original = cv2.resize(
                original, None, fx=output_scale, fy=output_scale, interpolation=cv2.INTER_AREA
            )
            scale_matrix = np.array(
                [[output_scale, 0, 0], [0, output_scale, 0], [0, 0, 1]], dtype=np.float64
            )
            inverse = scale_matrix @ inverse
        output_size = (original.shape[1], original.shape[0])
        overlay = cv2.warpPerspective(
            canonical[:, :, :3], inverse, output_size, flags=cv2.INTER_NEAREST
        )
        alpha = cv2.warpPerspective(
            canonical[:, :, 3], inverse, output_size, flags=cv2.INTER_NEAREST
        )
        mask = alpha > 0
        if np.any(mask):
            opacity = alpha[mask].astype(np.float32)[:, None] / 255.0
            original[mask] = (
                original[mask].astype(np.float32) * (1.0 - opacity)
                + overlay[mask].astype(np.float32) * opacity
            ).astype(np.uint8)
        original = _partial_status_banner(original, result)
        # Checked copies are review evidence; use a quality-controlled JPEG
        # so a 14 MB PNG does not dominate exports.
        ok, encoded = cv2.imencode(
            ".jpg", original, [cv2.IMWRITE_JPEG_QUALITY, 88, cv2.IMWRITE_JPEG_OPTIMIZE, 1]
        )
        if not ok:
            raise OSError("ไม่สามารถสร้างภาพผลตรวจได้")
        return encoded.tobytes()

    # Keep a clearly labelled fallback for records created before registration
    # provenance existed. It never mutates or overwrites the source image.
    width = max(original.shape[1], 1000)
    rows = (result["max"] + 3) // 4
    panel_height = 130 + rows * 32
    canvas = np.full((original.shape[0] + panel_height, width, 3), 255, np.uint8)
    canvas[: original.shape[0], : original.shape[1]] = original
    top = original.shape[0]
    cv2.putText(
        canvas,
        f"No. {result['student_number']}    Score {result['score']} / {result['max']}",
        (20, top + 38),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.85,
        (0, 0, 0),
        2,
    )
    cv2.putText(
        canvas,
        "Teacher reviewed | green OK | red X | orange blank | purple multiple | yellow boundary",
        (20, top + 72),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.5,
        (40, 40, 40),
        1,
    )
    for index, (answer, correct) in enumerate(zip(result["answers"], key["answers"], strict=True)):
        row, column = divmod(index, 4)
        if answer in {"blank", "multiple", "boundary_cross"}:
            status = {"blank": "blank", "multiple": "multiple", "boundary_cross": "boundary"}[
                answer
            ]
            text_color = CHECKED_COLORS[answer][:3]
        else:
            status = "OK" if score_answer(answer, correct) else "X"
            text_color = CHECKED_COLORS["correct" if status == "OK" else "incorrect"][:3]
        label = f"{index + 1}: {answer}/{correct} {status}"
        cv2.putText(
            canvas,
            label,
            (20 + column * (width // 4), top + 110 + row * 32),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.48,
            text_color,
            1,
        )
    canvas = _partial_status_banner(canvas, result)
    longest = max(canvas.shape[:2])
    if longest > 2400:
        scale = 2400 / longest
        canvas = cv2.resize(canvas, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    ok, encoded = cv2.imencode(
        ".jpg", canvas, [cv2.IMWRITE_JPEG_QUALITY, 88, cv2.IMWRITE_JPEG_OPTIMIZE, 1]
    )
    if not ok:
        raise OSError("ไม่สามารถสร้างภาพผลตรวจได้")
    return encoded.tobytes()


def export_results(
    flow: Workflow,
    exam_id: str,
    output_root: Path | None = None,
    room_id: str | None = None,
) -> Path:
    snapshot = flow.snapshot(exam_id, room_id=room_id)
    indicator_config = snapshot.get("assessment_indicators", {})
    if not indicator_config.get("valid", True):
        raise ValueError(
            "ตัวชี้วัดมีช่วงข้อไม่ถูกต้องตามจำนวนข้อปัจจุบัน กรุณาแก้หรือลบตัวชี้วัดก่อนออกผล"
        )
    indicators = indicator_config.get("items", [])
    snapshot_fingerprint = flow.snapshot_fingerprint(snapshot)
    importer = ImportService(flow.database)
    # Also verify the approved key's source evidence.
    key_source = next(
        record
        for record in importer.list_sources(exam_id)
        if record["id"] == snapshot["key"]["source_id"]
    )
    importer.verified_bytes(key_source)
    root_parent = (output_root or (flow.database.parent / "results")).expanduser().resolve()
    if root_parent.exists() and not root_parent.is_dir():
        raise OSError("ตำแหน่งบันทึกผลลัพธ์ไม่ใช่โฟลเดอร์")
    root_parent.mkdir(parents=True, exist_ok=True)
    if not os.access(root_parent, os.W_OK):
        raise OSError("ไม่มีสิทธิ์เขียนตำแหน่งบันทึกผลลัพธ์")
    room_label = snapshot["room"]["room_label"]
    folder_name = human_readable_name(snapshot["exam"], room_label=room_label)
    # Different labels can sanitize to the same cross-platform folder name
    # (for example, "ป.1/1" and "ป.1:1"). Keep those rooms deterministic and
    # separate while preserving the legacy folder name for non-colliding rooms.
    if any(
        other.id != snapshot["room"]["id"]
        and human_readable_name(snapshot["exam"], room_label=other.label) == folder_name
        for other in ExamStore(flow.database).list_rooms(exam_id)
    ):
        folder_name = f"{folder_name}-{snapshot['room']['id'][:8]}"
    root = root_parent / folder_name / "ผลการตรวจ"
    root.mkdir(parents=True, exist_ok=True)
    run_id = "run-" + str(uuid4())
    staging = Path(tempfile.mkdtemp(prefix=".staging-", dir=root))
    final = None
    published = False
    reserved = False
    try:
        (staging / "checked").mkdir()
        (staging / "_system").mkdir()
        book = Workbook()
        sheet = book.active
        sheet.title = "Scores"
        indicator_headers = [f"ตัวชี้วัด {item['identifier']}" for item in indicators]
        sheet.append(
            ["No.", *indicator_headers, "Score", "Max", "Status", "Source File", "Review Issues"]
        )
        used_names: set[str] = set()
        template_def = None
        try:
            template_def = load_exam_template_def(flow.database, exam_id)
        except Exception:
            template_def = None
        template_id = (
            template_def.template_id
            if template_def is not None
            else snapshot.get("exam", {}).get("template_id", "default-1")
        )

        for result in snapshot["results"]:
            data = importer.verified_bytes(result["source"])
            relative = f"checked/{checked_filename(result['student_number'], used_names)}"
            (staging / relative).write_bytes(
                checked_image(data, result, snapshot["key"], template_def=template_def)
            )
            result["checked_image"] = relative
            result["checked_provenance"] = {
                "source_sha256": result["source"]["sha256"],
                "detection_id": result.get("detection_id"),
                "review_id": result["review_id"],
                "partial_review_id": result.get("partial_review_id"),
                "identity_confirmed": result.get("identity_confirmed", True),
                "identity_origin": result.get("identity_origin", "teacher_confirmed"),
                "answer_provenance": result.get("answer_provenance", []),
                "review_issues": result.get("review_issues", []),
                "renderer_version": "checked-overlay-canonical-jpeg-v2-partial-review",
                "encoding": {"format": "JPEG", "quality": 88, "max_dimension": 2400},
                "color_legend": {
                    "correct": "green",
                    "incorrect": "red",
                    "blank": "orange",
                    "multiple": "purple",
                    "boundary_cross": "yellow",
                },
                "template_id": template_id,
                "template_version": snapshot.get("exam", {}).get("template_version", 1),
                "render_template": template_def.to_dict()
                if template_def is not None
                else template(),
            }
            sheet.append(
                [
                    result["student_number"],
                    *result.get("indicator_scores", []),
                    result["score"],
                    result["max"],
                    result["status"],
                    result["source"]["original_name"],
                    "; ".join(result.get("review_issues", [])),
                ]
            )
            # Force user-controlled strings to text, even if they begin with '='.
            text_columns = {
                1,
                4 + len(indicators),  # Status
                5 + len(indicators),  # Source File
                6 + len(indicators),  # Review Issues
            }
            for column in text_columns:
                sheet.cell(sheet.max_row, column).data_type = "s"
        sheet.freeze_panes = "A2"
        for attendance in snapshot.get("attendance", []):
            if attendance["status"] in {"absent", "excused"}:
                sheet.append(
                    [
                        str(attendance["student_number"]),
                        *([None] * len(indicators)),
                        None,
                        len(snapshot["key"]["answers"]),
                        "ขาดสอบ" if attendance["status"] == "absent" else "ลา / ได้รับยกเว้น",
                        None,
                        None,
                    ]
                )
        widths = [24, *([18] * len(indicators)), 12, 12, 24, 50, 50]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        info_sheet = book.create_sheet(title="Info")
        t_name = template_def.name if template_def is not None else template_id
        info_sheet.append(["หัวข้อ", "รายละเอียด"])
        info_sheet.append(["ชื่อข้อสอบ", snapshot["exam"]["name"]])
        info_sheet.append(["วิชา", snapshot["exam"]["subject"]])
        info_sheet.append(["ชั้น / ห้อง", f"{snapshot['exam']['grade']} / {room_label}"])
        info_sheet.append(["ปีการศึกษา", str(snapshot["exam"]["academic_year"])])
        info_sheet.append(["แม่แบบกระดาษคำตอบ", t_name])
        info_sheet.append(["จำนวนข้อ", str(snapshot["exam"]["question_count"])])
        for col_let, w in (("A", 22), ("B", 50)):
            info_sheet.column_dimensions[col_let].width = w
        book.active = sheet

        book.save(staging / "scores.xlsx")
        exam = snapshot["exam"]
        snapshot["run_id"] = run_id
        snapshot["app_version"] = __version__
        snapshot["export"] = {
            "human_readable_name": human_readable_name(exam, room_label=room_label),
            "folder_policy": "human-facing-v2",
            "template": template_def.to_dict() if template_def is not None else template(),
            "renderer_version": "checked-overlay-canonical-jpeg-v2-partial-review",
            "encoding": {"format": "JPEG", "quality": 88, "max_dimension": 2400},
            "color_legend": {
                "correct": "green",
                "incorrect": "red",
                "blank": "orange",
                "multiple": "purple",
                "boundary_cross": "yellow",
            },
            "snapshot_fingerprint": snapshot_fingerprint,
        }
        (staging / "_system" / "results.json").write_text(
            json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        hashes = {
            path.relative_to(staging).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in staging.rglob("*")
            if path.is_file()
        }
        (staging / "_system" / "manifest.json").write_text(
            json.dumps(hashes, indent=2), encoding="utf-8"
        )
        # Do not publish if a key/review or source set changed during rendering.
        current = flow.snapshot(exam_id, room_id=snapshot["room"]["id"])
        if flow.snapshot_fingerprint(current) != snapshot_fingerprint:
            raise ValueError("ข้อมูลเปลี่ยนระหว่างสร้างผล กรุณาสร้างผลใหม่")
        stamp = datetime.now().astimezone().strftime("%Y-%m-%d_%H%M%S")
        sequence = 1
        while True:
            final = root / (stamp if sequence == 1 else f"{stamp}_{sequence:02d}")
            try:
                final.mkdir()
                reserved = True
                break
            except FileExistsError:
                sequence += 1
        # Copy staging contents into final directory (Windows-safe alternative
        # to os.replace which fails on directories with WinError 5).
        for item in staging.rglob("*"):
            rel = item.relative_to(staging)
            dest = final / rel
            if item.is_dir():
                dest.mkdir(parents=True, exist_ok=True)
            else:
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(item, dest)
        published = True
        flow.record_export(
            exam_id,
            str(final),
            run_id,
            snapshot_fingerprint,
            room_id=snapshot["room"]["id"],
        )
        return final
    finally:
        if final is not None and reserved and not published:
            final.rmdir()
        if staging.exists():
            shutil.rmtree(staging)
